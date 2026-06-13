"""Document loader - PDF, DOCX, TXT, and image files.
Used for password-protected UiTM portals (Permata, uFuture, iStudent authenticated)
where documents are manually exported as per the report's data acquisition strategy.

Accepts either a local file path (for seed scripts) or raw bytes + filename
(for files streamed in from Supabase Storage / direct upload).
Images are processed with Claude Vision to extract text and describe visual content.
"""
import base64
from io import BytesIO
import pdfplumber
from docx import Document as DocxDocument
from datetime import datetime
from pathlib import Path
from typing import Optional, Union
from langchain_core.documents import Document

_IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".webp"}


def load_document(
    file_path: Optional[str] = None,
    portal_name: Optional[str] = None,
    source_url: Optional[str] = None,
    *,
    file_bytes: Optional[bytes] = None,
    filename: Optional[str] = None,
) -> list[Document]:
    """Extract text from a document, either from a local path or in-memory bytes.

    Args:
        file_path: Local path to the file (used by seed scripts).
        portal_name: e.g. 'Permata', 'uFuture'.
        source_url: Original URL or storage URL where the file lives.
        file_bytes: Raw file content (used when the file is in cloud storage).
        filename: Original filename (required when passing file_bytes).
    """
    if file_bytes is not None:
        if not filename:
            raise ValueError("filename is required when loading from bytes.")
        name = filename
        ext = Path(filename).suffix.lower()
        text = _extract_from_bytes(file_bytes, ext)
    elif file_path is not None:
        path = Path(file_path)
        name = path.name
        ext = path.suffix.lower()
        text = _extract_from_path(file_path, ext)
    else:
        raise ValueError("Either file_path or file_bytes must be provided.")

    if not text.strip():
        return []

    metadata = {
        "source": source_url or f"uploaded://{name}",
        "source_type": "document",
        "portal_name": portal_name or Path(name).stem,
        "title": name,
        "file_type": ext.lstrip("."),
        "scraped_at": datetime.utcnow().isoformat(),
    }
    return [Document(page_content=text, metadata=metadata)]


def _extract_from_path(file_path: str, ext: str) -> str:
    if ext == ".pdf":
        return _extract_pdf(file_path)
    if ext == ".docx":
        return _extract_docx(file_path)
    if ext in (".txt", ".md"):
        return Path(file_path).read_text(encoding="utf-8", errors="ignore")
    if ext in _IMAGE_EXTENSIONS:
        return _extract_image(Path(file_path).read_bytes(), ext)
    raise ValueError(f"Unsupported file type: {ext}")


def _extract_from_bytes(data: bytes, ext: str) -> str:
    if ext == ".pdf":
        return _extract_pdf(BytesIO(data))
    if ext == ".docx":
        return _extract_docx(BytesIO(data))
    if ext in (".txt", ".md"):
        return data.decode("utf-8", errors="ignore")
    if ext in _IMAGE_EXTENSIONS:
        return _extract_image(data, ext)
    raise ValueError(f"Unsupported file type: {ext}")


def _extract_pdf(source: Union[str, BytesIO]) -> str:
    chunks = []
    with pdfplumber.open(source) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            if page_text.strip():
                chunks.append(page_text)
    return "\n\n".join(chunks)


def _extract_docx(source: Union[str, BytesIO]) -> str:
    doc = DocxDocument(source)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)
    return "\n".join(paragraphs)


def _extract_image(data: bytes, ext: str) -> str:
    """Use Claude Vision to extract text and describe visual content from an image."""
    import anthropic
    from app.config import config

    media_types = {
        ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
        ".png": "image/png", ".gif": "image/gif", ".webp": "image/webp",
    }
    client = anthropic.Anthropic(api_key=config.ANTHROPIC_API_KEY)
    message = client.messages.create(
        model="claude-haiku-4-5-20251001",
        max_tokens=1024,
        messages=[{
            "role": "user",
            "content": [
                {
                    "type": "image",
                    "source": {
                        "type": "base64",
                        "media_type": media_types.get(ext, "image/jpeg"),
                        "data": base64.standard_b64encode(data).decode(),
                    },
                },
                {
                    "type": "text",
                    "text": (
                        "Extract all visible text from this image exactly as written. "
                        "Also describe any forms, tables, diagrams, or important visual elements. "
                        "This content is from a UiTM university portal — be thorough and accurate."
                    ),
                },
            ],
        }],
    )
    return message.content[0].text
